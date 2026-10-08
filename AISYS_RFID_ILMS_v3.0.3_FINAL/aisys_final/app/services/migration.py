from pathlib import Path
import csv
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from datetime import datetime
import shutil
from app.models.models import MigrationRun, MigrationRow, Book, SearchDocument

REQUIRED = {"accession_no", "isbn", "title", "author", "category", "reference_only"}


def read_rows(path: str):
    p = Path(path)
    if p.suffix.lower() == ".xlsx":
        ws = load_workbook(p, read_only=True, data_only=True).active
        rows = ws.iter_rows(values_only=True)
        try:
            headers = [str(x).strip() if x is not None else "" for x in next(rows)]
        except StopIteration:
            return
        missing = REQUIRED - set(headers)
        if missing:
            raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")
        for r in rows:
            yield dict(zip(headers, r))
    else:
        with open(p, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            headers = set(reader.fieldnames or [])
            missing = REQUIRED - headers
            if missing:
                raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")
            yield from reader


def validate_rows(db: Session, path: str):
    seen = set()
    results = []
    for i, row in enumerate(read_rows(path), start=2):
        accession = str(row.get("accession_no") or "").strip()
        title = str(row.get("title") or "").strip()
        status = "VALID"
        error = ""
        if not accession or not title:
            status, error = "INVALID", "accession_no and title are required"
        elif accession in seen or db.query(Book).filter(Book.accession_no == accession).first():
            status, error = "DUPLICATE", "duplicate accession_no"
        else:
            seen.add(accession)
        results.append((i, row, status, error))
    return results


def dry_run(db: Session, path: str):
    run = MigrationRun(source_name=Path(path).name, status="DRY_RUN")
    db.add(run)
    db.flush()
    try:
        results = validate_rows(db, path)
    except Exception as exc:
        run.status = "REJECTED"
        run.source_rows = 0
        run.invalid_rows = 1
        db.commit()
        return run
    for i, row, status, error in results:
        db.add(MigrationRow(run_id=run.id, row_number=i, accession_no=str(row.get("accession_no") or ""), status=status, error=error))
    run.source_rows = len(results)
    run.valid_rows = sum(s == "VALID" for _, _, s, _ in results)
    run.invalid_rows = sum(s == "INVALID" for _, _, s, _ in results)
    run.duplicate_rows = sum(s == "DUPLICATE" for _, _, s, _ in results)
    db.commit()
    return run


def import_file(db: Session, path: str):
    # Local/offline deployments automatically snapshot the SQLite database before import.
    # PostgreSQL deployments should use pg_dump as documented in the operations runbook.
    backup_path = None
    db_file = Path("aisys.db")
    if db_file.exists():
        backup_dir = Path("backups"); backup_dir.mkdir(exist_ok=True)
        backup_path = backup_dir / f"pre_migration_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
        shutil.copy2(db_file, backup_path)
    try:
        results = validate_rows(db, path)
    except Exception:
        run = MigrationRun(source_name=Path(path).name, status="REJECTED", source_rows=0, invalid_rows=1)
        db.add(run)
        db.commit()
        return run
    invalid = [r for r in results if r[2] != "VALID"]
    run = MigrationRun(
        source_name=Path(path).name,
        status="STARTED",
        source_rows=len(results),
        valid_rows=sum(r[2] == "VALID" for r in results),
        invalid_rows=sum(r[2] == "INVALID" for r in results),
        duplicate_rows=sum(r[2] == "DUPLICATE" for r in results),
    )
    db.add(run)
    db.flush()
    for i, row, status, error in results:
        db.add(MigrationRow(run_id=run.id, row_number=i, accession_no=str(row.get("accession_no") or ""), status=status, error=error))
    if invalid:
        run.status = "REJECTED"
        db.commit()
        return run
    try:
        for _, row, _, _ in results:
            ref = str(row.get("reference_only") or "").strip().lower() in {"true", "1", "yes", "y"}
            book = Book(
                accession_no=str(row.get("accession_no") or "").strip(),
                isbn=str(row.get("isbn") or "").strip(),
                title=str(row.get("title") or "").strip(),
                author=str(row.get("author") or "").strip(),
                category=str(row.get("category") or "General").strip(),
                reference_only=ref,
            )
            db.add(book)
            db.flush()
            db.add(SearchDocument(book_id=book.id, search_text=f"{book.accession_no} {book.title} {book.author} {book.category}"))
        run.migrated_rows = len(results)
        run.status = "COMPLETED"
        db.commit()
        return run
    except Exception:
        db.rollback()
        recovery = MigrationRun(
            source_name=Path(path).name,
            status="ROLLED_BACK",
            source_rows=len(results),
            valid_rows=len(results),
            migrated_rows=0,
        )
        db.add(recovery)
        db.commit()
        return recovery
